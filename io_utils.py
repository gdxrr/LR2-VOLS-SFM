"""
Утилиты загрузки/сохранения данных СФМ из Excel (.xlsx) и CSV.
"""

import csv
import numpy as np
from pathlib import Path
from typing import List

from calculations import SFMData

try:
    from openpyxl import Workbook, load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def _parse_float(val) -> float:
    if val is None:
        return 0.0
    try:
        return float(val)
    except (ValueError, TypeError):
        return 0.0


# ───────────────── CSV ─────────────────

def save_sfm_csv(path: str, sfm: SFMData):
    """Сохранить одну СФМ в CSV."""
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["name", sfm.name])
        w.writerow(["goal", sfm.goal])
        w.writerow(["date", sfm.date])
        w.writerow(["total_budget", sfm.total_budget])
        w.writerow(["budget_real", sfm.budget_real])
        w.writerow(["eval_period", sfm.eval_period])
        w.writerow(["period_real", sfm.period_real])

        w.writerow(["param_names"] + sfm.param_names)
        w.writerow(["param_units"] + sfm.param_units)

        header = []
        for pn in sfm.param_names:
            header.extend([f"{pn}_min", f"{pn}_max"])
        w.writerow(["# functions and TS"])

        row_idx = 0
        for fi, func_name in enumerate(sfm.functions):
            ts_list = sfm.ts_names[fi]
            for ti, ts_name in enumerate(ts_list):
                vals = sfm.values[row_idx].tolist()
                w.writerow([func_name, ts_name] + vals)
                row_idx += 1


def load_sfm_csv(path: str) -> SFMData:
    """Загрузить одну СФМ из CSV."""
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = list(csv.reader(f, delimiter=";"))

    meta = {}
    param_names = []
    param_units = []
    data_rows = []
    reading_data = False

    for row in reader:
        if not row:
            continue
        key = row[0].strip()

        if key == "# functions and TS":
            reading_data = True
            continue

        if reading_data:
            func_name = row[0]
            ts_name = row[1]
            vals = [_parse_float(v) for v in row[2:]]
            data_rows.append((func_name, ts_name, vals))
        elif key == "param_names":
            param_names = row[1:]
        elif key == "param_units":
            param_units = row[1:]
        else:
            meta[key] = row[1] if len(row) > 1 else ""

    functions = []
    ts_names = []
    current_func = None
    current_ts = []

    for func_name, ts_name, vals in data_rows:
        if func_name != current_func:
            if current_func is not None:
                functions.append(current_func)
                ts_names.append(current_ts)
            current_func = func_name
            current_ts = [ts_name]
        else:
            current_ts.append(ts_name)
    if current_func is not None:
        functions.append(current_func)
        ts_names.append(current_ts)

    num_cols = len(param_names) * 2
    values = np.zeros((len(data_rows), num_cols))
    for i, (_, _, vals) in enumerate(data_rows):
        for j, v in enumerate(vals[:num_cols]):
            values[i, j] = v

    return SFMData(
        name=meta.get("name", ""),
        goal=meta.get("goal", ""),
        date=meta.get("date", ""),
        functions=functions,
        ts_names=ts_names,
        param_names=param_names,
        param_units=param_units,
        values=values,
        total_budget=_parse_float(meta.get("total_budget", 0)),
        budget_real=_parse_float(meta.get("budget_real", 0)),
        eval_period=_parse_float(meta.get("eval_period", 0)),
        period_real=_parse_float(meta.get("period_real", 0)),
    )


# ───────────────── XLSX ─────────────────

def save_sfm_xlsx(path: str, sfm: SFMData):
    """Сохранить одну СФМ в Excel."""
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl не установлен")

    wb = Workbook()
    ws = wb.active
    ws.title = "СФМ"

    ws.append(["name", sfm.name])
    ws.append(["goal", sfm.goal])
    ws.append(["date", sfm.date])
    ws.append(["total_budget", sfm.total_budget])
    ws.append(["budget_real", sfm.budget_real])
    ws.append(["eval_period", sfm.eval_period])
    ws.append(["period_real", sfm.period_real])
    ws.append(["param_names"] + sfm.param_names)
    ws.append(["param_units"] + sfm.param_units)

    ws.append(["# functions and TS"])
    row_idx = 0
    for fi, func_name in enumerate(sfm.functions):
        for ti, ts_name in enumerate(sfm.ts_names[fi]):
            vals = sfm.values[row_idx].tolist()
            ws.append([func_name, ts_name] + vals)
            row_idx += 1

    wb.save(path)


def load_sfm_xlsx(path: str) -> SFMData:
    """Загрузить одну СФМ из Excel."""
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl не установлен")

    wb = load_workbook(path, data_only=True)
    ws = wb.active

    rows = []
    for row in ws.iter_rows(values_only=True):
        rows.append([str(c) if c is not None else "" for c in row])

    meta = {}
    param_names = []
    param_units = []
    data_rows = []
    reading_data = False

    for row in rows:
        if not row:
            continue
        key = row[0].strip()

        if key == "# functions and TS":
            reading_data = True
            continue

        if reading_data:
            func_name = row[0]
            ts_name = row[1]
            vals = [_parse_float(v) for v in row[2:]]
            data_rows.append((func_name, ts_name, vals))
        elif key == "param_names":
            param_names = [x for x in row[1:] if x]
        elif key == "param_units":
            param_units = [x for x in row[1:] if x]
        else:
            meta[key] = row[1] if len(row) > 1 else ""

    functions = []
    ts_names_list: List[List[str]] = []
    current_func = None
    current_ts: List[str] = []

    for func_name, ts_name, vals in data_rows:
        if func_name != current_func:
            if current_func is not None:
                functions.append(current_func)
                ts_names_list.append(current_ts)
            current_func = func_name
            current_ts = [ts_name]
        else:
            current_ts.append(ts_name)
    if current_func is not None:
        functions.append(current_func)
        ts_names_list.append(current_ts)

    num_cols = len(param_names) * 2
    values = np.zeros((len(data_rows), num_cols))
    for i, (_, _, vals) in enumerate(data_rows):
        for j, v in enumerate(vals[:num_cols]):
            values[i, j] = v

    return SFMData(
        name=meta.get("name", ""),
        goal=meta.get("goal", ""),
        date=meta.get("date", ""),
        functions=functions,
        ts_names=ts_names_list,
        param_names=param_names,
        param_units=param_units,
        values=values,
        total_budget=_parse_float(meta.get("total_budget", 0)),
        budget_real=_parse_float(meta.get("budget_real", 0)),
        eval_period=_parse_float(meta.get("eval_period", 0)),
        period_real=_parse_float(meta.get("period_real", 0)),
    )
